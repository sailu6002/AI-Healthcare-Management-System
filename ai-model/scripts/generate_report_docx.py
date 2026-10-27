"""Generate the Word (.docx) report from the Markdown source.

Reads documentation/AI-Model-Report.md and writes
documentation/AI-Model-Report.docx using a small Markdown subset:
headings, paragraphs, bullet/numbered lists, pipe tables, fenced code
blocks, horizontal rules, and inline **bold** / *italic* / `code`.

Run with the AI module's virtual environment:

    ai-model/venv/bin/python ai-model/scripts/generate_report_docx.py
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parents[2]
MD_PATH = REPO_ROOT / "documentation" / "AI-Model-Report.md"
DOCX_PATH = REPO_ROOT / "documentation" / "AI-Model-Report.docx"

MONO_FONT = "Consolas"
BODY_FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3B, 0x63)

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def _add_inline(paragraph, text: str) -> None:
    """Add runs to a paragraph, honouring **bold**, *italic*, `code`."""
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(10)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)


def _is_table_separator(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells) and bool(cells)


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.space_before = Pt(6)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = MONO_FONT
    run.font.size = Pt(8.5)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.paragraphs[0].text = ""
            _add_inline(cell.paragraphs[0], cell_text)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9.5)
                if r == 0:
                    run.bold = True
    doc.add_paragraph()


def _style_base(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    for level, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[level]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT


def build(md_text: str) -> Document:
    doc = Document()
    _style_base(doc)
    lines = md_text.splitlines()
    i = 0
    first_heading = True

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            _add_code_block(doc, block)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = _split_row(lines[i])
                if not _is_table_separator(cells):
                    rows.append(cells)
                i += 1
            _add_table(doc, rows)
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            if level == 1 and first_heading:
                heading = doc.add_heading(title, level=0)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_heading = False
            else:
                doc.add_heading(title, level=min(level, 4))
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)

        if bullet:
            _add_inline(doc.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered:
            _add_inline(doc.add_paragraph(style="List Number"), numbered.group(1))
        elif stripped:
            _add_inline(doc.add_paragraph(), stripped)
        i += 1

    doc.core_properties.title = "AI Risk Assessment Module - Mini Project Report"
    doc.core_properties.author = "AI-Based Smart Healthcare Management System"
    return doc


def main() -> int:
    if not MD_PATH.exists():
        print(f"Markdown source not found: {MD_PATH}", file=sys.stderr)
        return 1
    doc = build(MD_PATH.read_text(encoding="utf-8"))
    doc.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
